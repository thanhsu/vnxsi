"""Build the VNX.SI marketing brochure DOCX and its task-local visuals.

The brochure keeps its prose editable in Word. Illustrations are rendered to
PNG for stable inline placement; their SVG source companions are retained in
docs/marketing/assets for future editing.
"""

from __future__ import annotations

import math
import shutil
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "MARKETING.docx"
ASSET_DIR = ROOT / "docs" / "marketing" / "assets"

NAVY = "#0D1526"
BLUE = "#1D4ED8"
BLUE_SOFT = "#6B93FF"
BG = "#F4F5F7"
PAPER = "#FFFFFF"
TEXT_MUTED = "#5C6678"
LINE = "#D9DEE8"
PALE_BLUE = "#E9F0FF"
PALE_NAVY = "#E9EDF3"

FONT_DIR = Path(r"C:\Windows\Fonts")
FONT_REGULAR = FONT_DIR / "segoeui.ttf"
FONT_BOLD = FONT_DIR / "segoeuib.ttf"
FONT_LIGHT = FONT_DIR / "segoeuil.ttf"
DOC_FONT = "Segoe UI"


def hex_rgb(value: str) -> tuple[int, int, int]:
    value = value.lstrip("#")
    return tuple(int(value[i : i + 2], 16) for i in (0, 2, 4))


def font(size: int, bold: bool = False, light: bool = False) -> ImageFont.FreeTypeFont:
    path = FONT_BOLD if bold else FONT_LIGHT if light else FONT_REGULAR
    if not path.exists():
        path = FONT_DIR / "arialbd.ttf" if bold else FONT_DIR / "arial.ttf"
    return ImageFont.truetype(str(path), size)


def fit_text(draw: ImageDraw.ImageDraw, text: str, max_width: int, fnt: ImageFont.FreeTypeFont) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if current and draw.textlength(candidate, font=fnt) > max_width:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines


def text_block(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int],
    text: str,
    max_width: int,
    fnt: ImageFont.FreeTypeFont,
    fill: str,
    line_gap: int = 8,
) -> int:
    x, y = xy
    lines = fit_text(draw, text, max_width, fnt)
    line_height = fnt.size + line_gap
    draw.multiline_text((x, y), "\n".join(lines), font=fnt, fill=fill, spacing=line_gap)
    return y + line_height * len(lines)


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], fill: str, width: int = 5) -> None:
    draw.line((start[0], start[1], end[0], end[1]), fill=fill, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 15
    left = (end[0] - size * math.cos(angle - math.pi / 6), end[1] - size * math.sin(angle - math.pi / 6))
    right = (end[0] - size * math.cos(angle + math.pi / 6), end[1] - size * math.sin(angle + math.pi / 6))
    draw.polygon([end, left, right], fill=fill)


def circle_icon(draw: ImageDraw.ImageDraw, center: tuple[int, int], radius: int, fill: str, label: str, label_fill: str = PAPER) -> None:
    draw.ellipse((center[0] - radius, center[1] - radius, center[0] + radius, center[1] + radius), fill=fill)
    bbox = draw.textbbox((0, 0), label, font=font(radius // 2, bold=True))
    draw.text((center[0] - (bbox[2] - bbox[0]) / 2, center[1] - (bbox[3] - bbox[1]) / 2 - 2), label, font=font(radius // 2, bold=True), fill=label_fill)


def label(draw: ImageDraw.ImageDraw, xy: tuple[int, int], text: str, fill: str = BLUE) -> None:
    draw.text(xy, text.upper(), font=font(20, bold=True), fill=fill)


def round_card(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], fill: str = PAPER, outline: str = LINE, radius: int = 26, width: int = 2) -> None:
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def create_logo_png(path: Path, dark: bool = False) -> None:
    size = 512
    image = Image.new("RGBA", (size, size), (0, 0, 0, 0))
    draw = ImageDraw.Draw(image)
    stroke = hex_rgb("#E8EDF5" if dark else NAVY)
    fill = hex_rgb("#0A0F1C" if dark else PAPER)
    dot = hex_rgb("#6B93FF" if dark else BLUE)
    draw.line((112, 128, 256, 384, 400, 128), fill=stroke, width=40, joint="curve")
    for cx, cy in ((112, 128), (400, 128)):
        draw.ellipse((cx - 52, cy - 52, cx + 52, cy + 52), fill=fill, outline=stroke, width=30)
    draw.ellipse((256 - 64, 384 - 64, 256 + 64, 384 + 64), fill=dot)
    image.save(path, dpi=(300, 300))


def make_hero(path: Path) -> None:
    w, h = 1600, 760
    image = Image.new("RGB", (w, h), hex_rgb(BG))
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((720, 44, 1556, h - 44), radius=42, fill=hex_rgb(NAVY))
    label(draw, (70, 66), "Conceptual marketplace illustration")
    y = text_block(draw, (70, 135), "Find a starting point for the software you need.", 560, font(52, bold=True), NAVY, line_gap=7)
    text_block(draw, (70, y + 18), "A product, a builder, or a clear request can all lead to the right conversation.", 535, font(26), TEXT_MUTED, line_gap=9)
    center = (1138, 380)
    nodes = [
        ((850, 176), "PRODUCT", "A product that fits"),
        ((1410, 178), "BUILDER", "A person to meet"),
        ((850, 590), "NEED", "A request to shape"),
        ((1410, 590), "PATH", "Buy, customize, hire"),
    ]
    for (x, y2), title, sub in nodes:
        arrow(draw, (x, y2), center, hex_rgb("#9AB3FF"), width=5)
        draw.ellipse((x - 66, y2 - 66, x + 66, y2 + 66), fill=hex_rgb("#16213B"), outline=hex_rgb(BLUE_SOFT), width=4)
        bbox = draw.textbbox((0, 0), title, font=font(18, bold=True))
        draw.text((x - (bbox[2] - bbox[0]) / 2, y2 - 12), title, font=font(18, bold=True), fill=PAPER)
        bbox = draw.textbbox((0, 0), sub, font=font(17))
        draw.text((x - (bbox[2] - bbox[0]) / 2, y2 + 84), sub, font=font(17), fill=hex_rgb("#CBD5E9"))
    # Draw the central node last so connectors stay behind the disc and wordmark.
    draw.ellipse((center[0] - 112, center[1] - 112, center[0] + 112, center[1] + 112), fill=hex_rgb(BLUE))
    draw.ellipse((center[0] - 80, center[1] - 80, center[0] + 80, center[1] + 80), outline=hex_rgb("#93B0FF"), width=3)
    bbox = draw.textbbox((0, 0), "VNX.SI", font=font(35, bold=True))
    draw.text((center[0] - (bbox[2] - bbox[0]) / 2, center[1] - 22), "VNX.SI", font=font(35, bold=True), fill=PAPER)
    draw.text((70, h - 56), "VNX.SI / PRODUCT + PEOPLE", font=font(19, bold=True), fill=BLUE)
    image.save(path, dpi=(200, 200))


def make_routes(path: Path) -> None:
    w, h = 1600, 830
    image = Image.new("RGB", (w, h), hex_rgb(PAPER))
    draw = ImageDraw.Draw(image)
    label(draw, (62, 42), "Four practical starting points")
    routes = [
        ("BUY", "Choose a product that fits your work now.", "01", BLUE),
        ("CUSTOMIZE", "Start with something close and shape the next version.", "02", NAVY),
        ("HIRE", "Work with the builder behind a product.", "03", BLUE),
        ("BUILD SIMILAR", "Use a product as a reference for a different workflow.", "04", NAVY),
    ]
    boxes = [(62, 120, 752, 382), (848, 120, 1538, 382), (62, 452, 752, 714), (848, 452, 1538, 714)]
    for (title, desc, number, color), box in zip(routes, boxes):
        round_card(draw, box, fill=hex_rgb(BG), outline=hex_rgb(LINE), radius=30, width=3)
        x1, y1, x2, y2 = box
        draw.rounded_rectangle((x1, y1, x1 + 122, y1 + 12), radius=6, fill=hex_rgb(color))
        circle_icon(draw, (x1 + 74, y1 + 92), 39, color, number)
        draw.text((x1 + 136, y1 + 68), title, font=font(31, bold=True), fill=NAVY)
        text_block(draw, (x1 + 46, y1 + 164), desc, x2 - x1 - 92, font(25), TEXT_MUTED, line_gap=8)
        draw.rounded_rectangle((x1 + 46, y2 - 28, x2 - 46, y2 - 20), radius=4, fill=hex_rgb(BLUE_SOFT))
    image.save(path, dpi=(200, 200))


def make_journey(path: Path) -> None:
    w, h = 1600, 740
    image = Image.new("RGB", (w, h), hex_rgb(BG))
    draw = ImageDraw.Draw(image)
    label(draw, (70, 44), "From need to conversation")
    draw.text((70, 94), "Direct Inquiry or private request when needed.", font=font(37, bold=True), fill=NAVY)
    main_y = 292
    main = [(150, "01", "DISCOVER", "Find a product or builder"), (500, "02", "UNDERSTAND", "Read fit, context, and terms"), (850, "03", "INQUIRY", "Ask about a product or builder"), (1260, "04", "CONVERSATION", "Continue with the right person")]
    draw.line((150, main_y, 1260, main_y), fill=hex_rgb("#AFC2F7"), width=8)
    for idx, (x, num, title, desc) in enumerate(main):
        if idx < len(main) - 1:
            arrow(draw, (x + 50, main_y), (main[idx + 1][0] - 50, main_y), hex_rgb("#AFC2F7"), width=4)
        draw.ellipse((x - 50, main_y - 50, x + 50, main_y + 50), fill=hex_rgb(BLUE if idx in (0, 2) else NAVY), outline=hex_rgb(PAPER), width=5)
        bbox = draw.textbbox((0, 0), num, font=font(22, bold=True))
        draw.text((x - (bbox[2] - bbox[0]) / 2, main_y - 14), num, font=font(22, bold=True), fill=PAPER)
        tx = x - 116
        draw.text((tx, main_y + 83), title, font=font(18, bold=True), fill=NAVY)
        text_block(draw, (tx, main_y + 116), desc, 232, font(17), TEXT_MUTED, line_gap=5)
    # Optional private-request branch from understanding to proposal and back to conversation.
    branch_y = 530
    arrow(draw, (500, main_y + 52), (720, branch_y - 44), hex_rgb(BLUE_SOFT), width=4)
    branch = [(720, "A", "PRIVATE REQUEST", "Post when no listed product fits"), (1010, "B", "PROPOSALS", "Review builder responses")]
    draw.line((720, branch_y, 1010, branch_y), fill=hex_rgb("#AFC2F7"), width=6)
    for x, num, title, desc in branch:
        draw.ellipse((x - 44, branch_y - 44, x + 44, branch_y + 44), fill=hex_rgb(NAVY), outline=hex_rgb(PAPER), width=5)
        bbox = draw.textbbox((0, 0), num, font=font(21, bold=True))
        draw.text((x - (bbox[2] - bbox[0]) / 2, branch_y - 13), num, font=font(21, bold=True), fill=PAPER)
        tx = x - 118
        draw.text((tx, branch_y + 68), title, font=font(17, bold=True), fill=NAVY)
        text_block(draw, (tx, branch_y + 99), desc, 238, font(16), TEXT_MUTED, line_gap=4)
    # Route the return connector around the conversation label block.
    draw.line((1010, branch_y - 46, 1400, branch_y - 46), fill=hex_rgb(BLUE_SOFT), width=4)
    draw.line((1400, branch_y - 46, 1400, main_y), fill=hex_rgb(BLUE_SOFT), width=4)
    arrow(draw, (1400, main_y), (1310, main_y), hex_rgb(BLUE_SOFT), width=4)
    draw.rounded_rectangle((72, h - 74, 1528, h - 25), radius=16, fill=hex_rgb(PAPER), outline=hex_rgb(LINE), width=2)
    draw.text((98, h - 60), "Private requests invite up to five builders. Email stays hidden through the flow.", font=font(19, bold=True), fill=BLUE)
    image.save(path, dpi=(200, 200))


def make_builder(path: Path) -> None:
    w, h = 1600, 790
    image = Image.new("RGB", (w, h), hex_rgb(NAVY))
    draw = ImageDraw.Draw(image)
    label(draw, (72, 54), "A durable presence for useful work", fill=BLUE_SOFT)
    text_block(draw, (72, 115), "Build once. Present it clearly. Stay available for the next conversation.", 630, font(47, bold=True), PAPER, line_gap=8)
    draw.text((72, 350), "BUILDER WORK LOOP", font=font(19, bold=True), fill=hex_rgb("#9AB3FF"))
    center = (1110, 378)
    stages = [(840, 176, "PUBLISH", "Profile + product"), (1378, 176, "SHOW PROOF", "Demo + badges"), (840, 590, "ANSWER", "Inquiry + request"), (1378, 590, "ADAPT", "Customize + hire")]
    for x, y, title, sub in stages:
        arrow(draw, (x, y), center, hex_rgb("#617AB5"), width=5)
        draw.ellipse((x - 70, y - 70, x + 70, y + 70), fill=hex_rgb("#17233E"), outline=hex_rgb("#6B93FF"), width=4)
        bbox = draw.textbbox((0, 0), title, font=font(17, bold=True))
        draw.text((x - (bbox[2] - bbox[0]) / 2, y - 9), title, font=font(17, bold=True), fill=PAPER)
        bbox = draw.textbbox((0, 0), sub, font=font(17))
        draw.text((x - (bbox[2] - bbox[0]) / 2, y + 91), sub, font=font(17), fill=hex_rgb("#CBD5E9"))
    # Draw the central node last so the work-loop connectors stay behind it.
    draw.ellipse((center[0] - 98, center[1] - 98, center[0] + 98, center[1] + 98), fill=hex_rgb(BLUE))
    bbox = draw.textbbox((0, 0), "BUILDER", font=font(24, bold=True))
    draw.text((center[0] - (bbox[2] - bbox[0]) / 2, center[1] - 14), "BUILDER", font=font(24, bold=True), fill=PAPER)
    draw.rounded_rectangle((72, 600, 650, 705), radius=22, fill=hex_rgb("#16213B"), outline=hex_rgb("#34486D"), width=2)
    draw.text((102, 625), "Model agnostic", font=font(23, bold=True), fill=BLUE_SOFT)
    draw.text((102, 659), "Use the AI tools that suit your process.", font=font(20), fill=hex_rgb("#CBD5E9"))
    image.save(path, dpi=(200, 200))


def make_trust(path: Path) -> None:
    w, h = 1600, 720
    image = Image.new("RGB", (w, h), hex_rgb(PAPER))
    draw = ImageDraw.Draw(image)
    label(draw, (70, 48), "Evidence-based trust signals")
    draw.text((70, 98), "Each badge says what was checked.", font=font(38, bold=True), fill=NAVY)
    badges = [(330, 330, "LISTED", "Clear listing + working links", BLUE), (800, 330, "DEMO VERIFIED", "Demo opened over HTTPS", NAVY), (1270, 330, "IN PRODUCTION", "Evidence of real use", BLUE)]
    for x, y, title, desc, color in badges:
        draw.ellipse((x - 142, y - 142, x + 142, y + 142), fill=hex_rgb(BG), outline=hex_rgb(color), width=10)
        draw.ellipse((x - 112, y - 112, x + 112, y + 112), outline=hex_rgb("#B8C8EA"), width=3)
        draw.line((x - 46, y + 2, x - 12, y + 38, x + 62, y - 48), fill=hex_rgb(color), width=13, joint="curve")
        bbox = draw.textbbox((0, 0), title, font=font(22, bold=True))
        draw.text((x - (bbox[2] - bbox[0]) / 2, y + 166), title, font=font(22, bold=True), fill=NAVY)
        text_block(draw, (x - 130, y + 202), desc, 260, font(18), TEXT_MUTED, line_gap=5)
    draw.rounded_rectangle((70, h - 94, 1530, h - 40), radius=18, fill=hex_rgb(PALE_BLUE))
    draw.text((98, h - 76), "These signals describe the named checks; they do not replace your own review.", font=font(20, bold=True), fill=NAVY)
    image.save(path, dpi=(200, 200))


def make_roadmap(path: Path) -> None:
    w, h = 1600, 730
    image = Image.new("RGB", (w, h), hex_rgb(BG))
    draw = ImageDraw.Draw(image)
    label(draw, (74, 50), "Building in stages")
    draw.text((74, 98), "Make supply discoverable first. Add assistance carefully.", font=font(37, bold=True), fill=NAVY)
    y = 352
    draw.line((160, y, 1440, y), fill=hex_rgb("#9EB3E7"), width=8)
    points = [(230, "NOW", "Supply + conversation", "Profiles, products, catalogue, Inquiry, request", BLUE), (800, "NEXT", "AI discovery", "Plain-language idea to structured brief", NAVY), (1370, "LATER", "Project workflows", "Milestones, payments, reviews, maintenance", NAVY)]
    for x, stage, title, desc, color in points:
        draw.ellipse((x - 56, y - 56, x + 56, y + 56), fill=hex_rgb(color), outline=hex_rgb(PAPER), width=6)
        bbox = draw.textbbox((0, 0), stage, font=font(18, bold=True))
        draw.text((x - (bbox[2] - bbox[0]) / 2, y - 12), stage, font=font(18, bold=True), fill=PAPER)
        draw.text((x - 145, y + 90), title, font=font(23, bold=True), fill=NAVY)
        text_block(draw, (x - 145, y + 128), desc, 290, font(18), TEXT_MUTED, line_gap=6)
    draw.rounded_rectangle((70, h - 100, 1530, h - 45), radius=18, fill=hex_rgb(PAPER), outline=hex_rgb(LINE), width=2)
    draw.text((98, h - 82), "Roadmap directions are future plans, not current product promises.", font=font(20, bold=True), fill=BLUE)
    image.save(path, dpi=(200, 200))


def write_svg_sources() -> None:
    """Keep lightweight editable vector companions for the generated PNGs."""
    svg_sources = {
        "hero.svg": f'''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1600 760"><rect width="1600" height="760" fill="{BG}"/><rect x="720" y="44" width="836" height="672" rx="42" fill="{NAVY}"/><g fill="none" stroke="#9AB3FF" stroke-width="5"><path d="M850 176 L1138 380 L1410 178"/><path d="M850 590 L1138 380 L1410 590"/></g><g fill="#16213B" stroke="#6B93FF" stroke-width="4"><circle cx="850" cy="176" r="66"/><circle cx="1410" cy="178" r="66"/><circle cx="850" cy="590" r="66"/><circle cx="1410" cy="590" r="66"/></g><circle cx="1138" cy="380" r="112" fill="{BLUE}"/><circle cx="1138" cy="380" r="80" fill="none" stroke="#93B0FF" stroke-width="3"/><text x="1138" y="392" text-anchor="middle" font-family="Arial" font-size="35" font-weight="700" fill="white">VNX.SI</text><g font-family="Arial" font-size="18" font-weight="700" text-anchor="middle" fill="white"><text x="850" y="182">PRODUCT</text><text x="1410" y="184">BUILDER</text><text x="850" y="596">NEED</text><text x="1410" y="596">PATH</text></g></svg>''',
        "routes.svg": f'''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1600 830"><rect width="1600" height="830" fill="white"/><g fill="{BG}" stroke="{LINE}" stroke-width="3"><rect x="62" y="120" width="690" height="262" rx="30"/><rect x="848" y="120" width="690" height="262" rx="30"/><rect x="62" y="452" width="690" height="262" rx="30"/><rect x="848" y="452" width="690" height="262" rx="30"/></g><g fill="{BLUE}"><rect x="62" y="120" width="122" height="12" rx="6"/><rect x="848" y="120" width="122" height="12" rx="6"/></g><g fill="{NAVY}"><rect x="62" y="452" width="122" height="12" rx="6"/><rect x="848" y="452" width="122" height="12" rx="6"/></g><g font-family="Arial" font-size="31" font-weight="700" fill="{NAVY}"><text x="198" y="219">BUY</text><text x="984" y="219">CUSTOMIZE</text><text x="198" y="551">HIRE</text><text x="984" y="551">BUILD SIMILAR</text></g></svg>''',
        "journey.svg": f'''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1600 740"><rect width="1600" height="740" fill="{BG}"/><line x1="150" y1="292" x2="1260" y2="292" stroke="#AFC2F7" stroke-width="8"/><path d="M500 344 L720 486" fill="none" stroke="#6B93FF" stroke-width="4"/><line x1="720" y1="530" x2="1010" y2="530" stroke="#AFC2F7" stroke-width="6"/><path d="M1010 486 H1400 V292 H1310" fill="none" stroke="#6B93FF" stroke-width="4"/><g fill="{BLUE}"><circle cx="150" cy="292" r="50"/><circle cx="850" cy="292" r="50"/></g><g fill="{NAVY}"><circle cx="500" cy="292" r="50"/><circle cx="1260" cy="292" r="50"/><circle cx="720" cy="530" r="44"/><circle cx="1010" cy="530" r="44"/></g><g font-family="Arial" font-size="18" font-weight="700" text-anchor="middle" fill="white"><text x="150" y="298">01</text><text x="500" y="298">02</text><text x="850" y="298">03</text><text x="1260" y="298">04</text><text x="720" y="536">A</text><text x="1010" y="536">B</text></g></svg>''',
        "builder.svg": f'''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1600 790"><rect width="1600" height="790" fill="{NAVY}"/><g fill="none" stroke="#617AB5" stroke-width="5"><path d="M840 176 L1110 378 L1378 176"/><path d="M840 590 L1110 378 L1378 590"/></g><g fill="#17233E" stroke="#6B93FF" stroke-width="4"><circle cx="840" cy="176" r="70"/><circle cx="1378" cy="176" r="70"/><circle cx="840" cy="590" r="70"/><circle cx="1378" cy="590" r="70"/></g><circle cx="1110" cy="378" r="98" fill="{BLUE}"/><text x="1110" y="386" text-anchor="middle" font-family="Arial" font-size="24" font-weight="700" fill="white">BUILDER</text></svg>''',
        "trust.svg": f'''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1600 720"><rect width="1600" height="720" fill="white"/><g fill="{BG}" stroke="{BLUE}" stroke-width="10"><circle cx="330" cy="330" r="142"/><circle cx="800" cy="330" r="142"/><circle cx="1270" cy="330" r="142"/></g><g fill="none" stroke="{BLUE}" stroke-width="13" stroke-linecap="round" stroke-linejoin="round"><path d="M284 332 l34 36 l74 -84"/><path d="M754 332 l34 36 l74 -84"/><path d="M1224 332 l34 36 l74 -84"/></g></svg>''',
        "roadmap.svg": f'''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1600 730"><rect width="1600" height="730" fill="{BG}"/><line x1="160" y1="352" x2="1440" y2="352" stroke="#9EB3E7" stroke-width="8"/><g stroke="white" stroke-width="6"><circle cx="230" cy="352" r="56" fill="{BLUE}"/><circle cx="800" cy="352" r="56" fill="{NAVY}"/><circle cx="1370" cy="352" r="56" fill="{NAVY}"/></g></svg>''',
    }
    for filename, content in svg_sources.items():
        (ASSET_DIR / filename).write_text(content, encoding="utf-8")


def create_assets() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    write_svg_sources()
    source_logo = ROOT / "apps" / "web" / "public" / "assets" / "brand" / "vnxsi-mark.svg"
    shutil.copy2(source_logo, ASSET_DIR / "vnxsi-mark.svg")
    create_logo_png(ASSET_DIR / "vnxsi-mark.png")
    paths = {
        "logo": ASSET_DIR / "vnxsi-mark.png",
        "hero": ASSET_DIR / "hero.png",
        "routes": ASSET_DIR / "routes.png",
        "journey": ASSET_DIR / "journey.png",
        "builder": ASSET_DIR / "builder.png",
        "trust": ASSET_DIR / "trust.png",
        "roadmap": ASSET_DIR / "roadmap.png",
    }
    make_hero(paths["hero"])
    make_routes(paths["routes"])
    make_journey(paths["journey"])
    make_builder(paths["builder"])
    make_trust(paths["trust"])
    make_roadmap(paths["roadmap"])
    return paths


def set_font(run, name: str, size: float, color: str = NAVY, bold: bool | None = None, italic: bool = False) -> None:
    # The brand fonts are the design reference. Segoe UI is the installed
    # Office-safe fallback used in the artifact so the brochure remains a
    # clean sans serif when the brand fonts are not installed on the reader's
    # machine.
    output_name = DOC_FONT if name in {"Space Grotesk", "Be Vietnam Pro"} else name
    run.font.name = output_name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), output_name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), output_name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), output_name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(*hex_rgb(color))
    if bold is not None:
        run.bold = bold
    run.italic = italic


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill.lstrip("#"))


def set_cell_border(cell, color: str = LINE, size: str = "8") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color.lstrip("#"))


def set_cell_margins(cell, top: int = 120, start: int = 160, bottom: int = 120, end: int = 160) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    node = p_pr.find(qn("w:keepNext"))
    if node is None:
        node = OxmlElement("w:keepNext")
        p_pr.append(node)


def keep_together(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    node = p_pr.find(qn("w:keepLines"))
    if node is None:
        node = OxmlElement("w:keepLines")
        p_pr.append(node)


def remove_paragraph_borders(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def remove_style_borders(style) -> None:
    p_pr = style._element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE.lstrip("#"))
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_image_alt(run, description: str) -> None:
    doc_prs = run._r.xpath(".//wp:docPr")
    if doc_prs:
        doc_prs[0].set("descr", description)
        doc_prs[0].set("title", description)


def add_figure(doc: Document, image_path: Path, width: float, alt: str, caption: str, dark: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width))
    set_image_alt(run, alt)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    r = cap.add_run(caption)
    set_font(r, "Be Vietnam Pro", 8.3, "#667085", italic=True)


def add_body(doc: Document, text: str, after: float = 6, before: float = 0) -> None:
    p = doc.add_paragraph(style="Body Text")
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_font(r, "Be Vietnam Pro", 10.4, NAVY)


def add_body_with_lead(doc: Document, lead: str, text: str, after: float = 6) -> None:
    p = doc.add_paragraph(style="Body Text")
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(lead)
    set_font(r, "Be Vietnam Pro", 10.4, NAVY, bold=True)
    r = p.add_run(text)
    set_font(r, "Be Vietnam Pro", 10.4, NAVY)


def add_kicker(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    keep_with_next(p)
    r = p.add_run(text.upper())
    set_font(r, "Be Vietnam Pro", 8.5, BLUE, bold=True)


def add_h1(doc: Document, text: str, after: float = 8) -> None:
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(after)
    keep_with_next(p)
    r = p.add_run(text)
    set_font(r, "Space Grotesk", 18.5, "#000000", bold=True)


def add_h2(doc: Document, text: str, after: float = 3) -> None:
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(after)
    keep_with_next(p)
    r = p.add_run(text)
    set_font(r, "Space Grotesk", 12.4, "#000000", bold=True)


def add_numbered_step(doc: Document, number: str, title: str, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.05)
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(f"{number}  ")
    set_font(r, "Space Grotesk", 10.5, BLUE, bold=True)
    r = p.add_run(title)
    set_font(r, "Be Vietnam Pro", 10.3, NAVY, bold=True)
    r = p.add_run(f"  {text}")
    set_font(r, "Be Vietnam Pro", 10.3, NAVY)


def add_page_marker(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text.upper())
    set_font(r, "Be Vietnam Pro", 8.3, TEXT_MUTED, bold=True)


def add_logo_line(doc: Document, logo_path: Path) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run()
    r.add_picture(str(logo_path), width=Inches(0.42))
    set_image_alt(r, "VNX.SI connected nodes logo")
    r = p.add_run("   VNX.SI")
    set_font(r, "Space Grotesk", 15, NAVY, bold=True)


def add_footer(section) -> None:
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.75), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run("VNX.SI")
    set_font(r, "Space Grotesk", 8, NAVY, bold=True)
    r = p.add_run("\tMarketing overview | October 2026")
    set_font(r, "Be Vietnam Pro", 8, TEXT_MUTED)


def add_badge_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(1.55), Inches(5.25)]
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, text in enumerate(("Badge", "What the check means")):
        cell = hdr.cells[idx]
        cell.width = widths[idx]
        shade_cell(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_font(r, "Be Vietnam Pro", 9.3, PAPER, bold=True)
    rows = [
        ("Listed", "The VNX.SI team reviewed the listing for a clear description, pricing or a clear way to request pricing, and working links before publication."),
        ("Demo verified", "The team opened the demo and saw it work over HTTPS at the time of review. A changed demo URL revokes the badge."),
        ("In production", "The builder provided evidence of real customers using the product, and the team checked that evidence."),
    ]
    for i, (name, desc) in enumerate(rows):
        cells = table.add_row().cells
        set_cant_split(table.rows[-1])
        for idx, text in enumerate((name, desc)):
            cell = cells[idx]
            cell.width = widths[idx]
            shade_cell(cell, PAPER if i % 2 == 0 else BG)
            set_cell_border(cell)
            set_cell_margins(cell, top=130, bottom=130)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(text)
            set_font(r, "Be Vietnam Pro", 9.1 if idx else 9.3, NAVY, bold=(idx == 0))


def add_two_column_table(doc: Document, left_title: str, left_text: str, right_title: str, right_text: str) -> None:
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    set_cant_split(table.rows[0])
    set_cant_split(table.rows[1])
    for idx, title in enumerate((left_title, right_title)):
        cell = table.rows[0].cells[idx]
        cell.width = Inches(3.38)
        shade_cell(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell, top=120, bottom=120, start=190, end=190)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(title)
        set_font(r, "Space Grotesk", 11.2, PAPER, bold=True)
    for idx, text in enumerate((left_text, right_text)):
        cell = table.rows[1].cells[idx]
        cell.width = Inches(3.38)
        shade_cell(cell, PALE_BLUE if idx == 0 else PALE_NAVY)
        set_cell_border(cell)
        set_cell_margins(cell, top=150, bottom=150, start=190, end=190)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(text)
        set_font(r, "Be Vietnam Pro", 9.8, NAVY)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = DOC_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), DOC_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), DOC_FONT)
    normal.font.size = Pt(10.4)
    normal.font.color.rgb = RGBColor(*hex_rgb(NAVY))
    for style_name, size in (("Title", 28), ("Heading 1", 18.5), ("Heading 2", 12.4)):
        style = styles[style_name]
        style.font.name = DOC_FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), DOC_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), DOC_FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        remove_style_borders(style)
    if "Body Text" not in styles:
        styles.add_style("Body Text", WD_STYLE_TYPE.PARAGRAPH)
    body = styles["Body Text"]
    body.base_style = styles["Normal"]
    body.font.name = DOC_FONT
    body._element.rPr.rFonts.set(qn("w:ascii"), DOC_FONT)
    body._element.rPr.rFonts.set(qn("w:hAnsi"), DOC_FONT)
    body.font.size = Pt(10.4)
    body.font.color.rgb = RGBColor(*hex_rgb(NAVY))


def build_docx(assets: dict[str, Path]) -> None:
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.67)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.28)
    configure_styles(doc)
    add_footer(section)

    props = doc.core_properties
    props.title = "VNX.SI Marketing Overview"
    props.subject = "VNX.SI marketplace introduction"
    props.author = "VNX.SI"
    props.last_modified_by = "VNX.SI"
    props.comments = ""

    # Page 1: opening.
    add_logo_line(doc, assets["logo"])
    add_page_marker(doc, "Marketing overview / 01")
    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.keep_with_next = True
    remove_paragraph_borders(title)
    r = title.add_run("VNX.SI Marketing Overview")
    set_font(r, "Space Grotesk", 28, "#000000", bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(13)
    r = subtitle.add_run("Software you need. Built by people you can meet.")
    set_font(r, "Space Grotesk", 14.5, "#000000", bold=True)
    add_figure(
        doc,
        assets["hero"],
        6.82,
        "Conceptual illustration of VNX.SI connecting a client need with products, builders, and practical next steps.",
        "Conceptual marketplace illustration. VNX.SI brings products, builders, and needs into one clear starting point.",
    )
    add_body(doc, "VNX.SI is a marketplace for AI-built products and the builders behind them. People who need software can start with a product that already fits, ask for a tailored version, work directly with a builder, or describe what they need when the catalogue does not have the answer.", after=6)
    add_body_with_lead(doc, "Current scope - ", "As of October 2026, VNX.SI focuses on builder profiles, product listings, catalogue discovery, Inquiry, and Post a request. The platform does not process payments; clients and builders agree final scope, price, licence, and delivery directly.", after=0)

    # Page 2: routes.
    doc.add_page_break()
    add_page_marker(doc, "Start with the outcome / 02")
    add_h1(doc, "Choose the route that fits")
    add_body(doc, "You do not need to begin by choosing a framework or a development process. Begin with the work you want software to improve, then choose the route that matches the distance between your need and the available supply.", after=8)
    add_figure(doc, assets["routes"], 6.82, "Four conceptual paths on VNX.SI: Buy, Customize, Hire the builder, and Build Similar.", "Conceptual route map. The four paths describe how a client can move from a need to a useful conversation.")
    add_h2(doc, "Buy")
    add_body(doc, "Inspect the problem statement, features, demo, pricing tiers, licence, support policy, and builder. Send a Buy Inquiry when you want to ask a question or discuss next steps.", after=3)
    add_h2(doc, "Customize")
    add_body(doc, "If a product is close but needs changes, choose Customize when the builder offers it. Share context, budget range, and deadline so the conversation starts with a concrete brief.", after=3)
    add_h2(doc, "Hire the builder")
    add_body(doc, "Review the builder's profile, skills, portfolio, and availability, then send a Hire Inquiry for a separate project.", after=3)
    add_h2(doc, "Build Similar")
    add_body(doc, "Use a product as a useful starting point for a different business or workflow. Explain what you want to keep and what needs to change.", after=0)

    # Page 3: journey.
    doc.add_page_break()
    add_page_marker(doc, "A clearer route from idea to conversation / 03")
    add_h1(doc, "From need to conversation")
    add_body(doc, "VNX.SI keeps the first step concrete. A client can discover supply, understand what a product or builder offers, and choose a direct Inquiry or a private request without handing over an email address through the flow.", after=8)
    add_figure(doc, assets["journey"], 6.82, "Conceptual journey showing a direct Inquiry path and an optional private request path that meet at conversation.", "The paths meet at conversation: use a direct Inquiry for a product or builder, or a private request when the catalogue is not enough.")
    add_numbered_step(doc, "01", "Discover.", "Search the catalogue by category, delivery model, starting price, language, or verification badge. Browse builders by skill, language, country, and availability.")
    add_numbered_step(doc, "02", "Understand.", "Read the product page for its problem, intended users, features, technology context, pricing, licence, customization, support, and builder.")
    add_numbered_step(doc, "03", "Choose a route.", "Send a direct Inquiry when a product or builder is close, or post a private request when the catalogue does not have the right fit.")
    add_numbered_step(doc, "04", "Open an Inquiry.", "Share a message, budget range, and optional deadline. The conversation continues in the VNX.SI web inbox.")
    add_numbered_step(doc, "05", "Review proposals when needed.", "For a private request, the VNX.SI team can invite up to five builders to respond with proposals; review them in your account.")
    add_numbered_step(doc, "06", "Continue.", "Select a direction and keep the discussion focused on scope, price, licence, and delivery.")
    add_body_with_lead(doc, "Privacy by default - ", "Signed-out clients confirm their email before an Inquiry or request is opened for a builder to see. Builders do not receive a client's email through these flows.", after=0)

    # Page 4: builders.
    doc.add_page_break()
    add_page_marker(doc, "For builders / 04")
    add_h1(doc, "Turn useful work into a durable presence")
    add_body(doc, "AI can make building faster. VNX.SI gives what you build a place to be found, understood, and connected to the next conversation. A product can represent a SaaS, source-code offering, or service package.", after=8)
    add_figure(doc, assets["builder"], 6.82, "Conceptual builder work loop linking publishing, proof, response, and adaptation.", "Conceptual builder work loop. A clear presence lets one useful piece of work support several kinds of conversation.")
    add_two_column_table(
        doc,
        "Publish the work",
        "Create a profile with skills, working languages, availability, and portfolio. Publish a product with its problem, target users, features, demo, pricing, licence, customization options, and support policy.",
        "Continue the conversation",
        "Receive Inquiries from people who want to buy, customize, hire, or build something similar. Respond to private requests with an approach, price or price range, and expected timeline.",
    )
    add_h2(doc, "A channel for more than one project")
    add_body(doc, "The catalogue lets builders present products instead of explaining the same work from scratch. During the founding phase, listings are intended to be free while the initial supply is built.", after=5)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Become a builder at ")
    set_font(r, "Be Vietnam Pro", 10.4, NAVY, bold=True)
    add_hyperlink(p, "vnx.si/for-builders", "https://vnx.si/for-builders")

    # Page 5: trust.
    doc.add_page_break()
    add_page_marker(doc, "Trust you can read / 05")
    add_h1(doc, "Trust you can read")
    add_body(doc, "VNX.SI uses evidence-based badges instead of star ratings it cannot substantiate. Each active badge appears with a verification date and names the check that supports it.", after=8)
    add_figure(doc, assets["trust"], 6.82, "Three evidence-based VNX.SI badges: Listed, Demo verified, and In production.", "Evidence-based badges. Each signal describes the check it names.")
    add_badge_table(doc)
    add_body(doc, "These signals do not replace your own review of a product's fit, security, performance, licence, or commercial terms. A badge is a readable starting point, not a guarantee.", after=6, before=10)
    add_h2(doc, "Organic ranking has a clear line")
    add_body(doc, "Organic ranking is never for sale. Product and builder lists cannot be moved by payment, partner commission, or monetization data. If VNX.SI introduces sponsored placement in a later phase, it will be a separate, clearly labelled area and will not change organic order.", after=5)
    add_body(doc, "The same separation guides partner work. Future partner links may be disclosed and measured through controlled routes; a click or commission cannot become a quality signal. Partnership conversations can start at contact@vnx.si.", after=0)

    # Page 6: roadmap and CTA.
    doc.add_page_break()
    add_page_marker(doc, "What comes next / 06")
    add_h1(doc, "Build in stages")
    add_body(doc, "The first phase is about making supply discoverable and conversations useful enough to support the next one. Later phases are planned directions, evaluated before they become product promises.", after=8)
    add_figure(doc, assets["roadmap"], 6.82, "Roadmap timeline distinguishing current supply and conversation features from future discovery and project workflows.", "Roadmap direction. The future stages are clearly separated from the current product scope.")
    add_two_column_table(
        doc,
        "Available now",
        "Builder profiles, product listings, catalogue discovery, Inquiry, and Post a request. The interface supports English, Vietnamese, Simplified Chinese, and Traditional Chinese.",
        "Planned directions",
        "AI Discovery, recommendations grounded in real products and builders, assisted matching and estimates, and project, milestone, payment, payout, review, and maintenance workflows.",
    )
    add_h2(doc, "Start with VNX.SI")
    add_body(doc, "Have a software idea but unsure whether to buy, customize, or build? Start with real products and builders, then turn the need into a clear conversation.", after=5)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("Explore products and builders at ")
    set_font(r, "Be Vietnam Pro", 10.8, NAVY, bold=True)
    add_hyperlink(p, "vnx.si", "https://vnx.si")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Press, partnerships, and general questions: ")
    set_font(r, "Be Vietnam Pro", 10.2, TEXT_MUTED)
    add_hyperlink(p, "contact@vnx.si", "mailto:contact@vnx.si")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


def main() -> None:
    assets = create_assets()
    build_docx(assets)
    print(f"created {OUT}")
    print(f"assets {ASSET_DIR}")


if __name__ == "__main__":
    main()
